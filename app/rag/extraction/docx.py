"""DOCX text extraction strategy."""
from __future__ import annotations

from io import BytesIO
from typing import ClassVar

from docx import Document

from app.rag.extraction.base import TextExtractor


class DocxExtractor(TextExtractor):
    """Extract text from Word ``.docx`` files using :mod:`python-docx`."""

    extensions: ClassVar[frozenset[str]] = frozenset({".docx"})

    def extract(self, data: bytes) -> str:
        document = Document(BytesIO(data))

        parts: list[str] = [
            paragraph.text for paragraph in document.paragraphs
        ]

        # Include text from tables, which are not part of ``paragraphs``.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells if cell.text]
                if cells:
                    parts.append("\t".join(cells))

        return "\n".join(parts)
